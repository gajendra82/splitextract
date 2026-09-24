"""Format-independent Excel stock-statement extraction."""

from __future__ import annotations

import io
import unittest
from datetime import datetime
from pathlib import Path

from services.sales_statement_extractor import (
    _to_float,
    _xls_norm_header,
    extract_sales_statement,
)


SAMPLE_XLS = Path(
    r"C:\Users\gajen\Downloads\testing_files\0000700248_2026_08_ZL_18_245_05092026055923.XLS"
)
KNOWN_XLSX = (
    Path(__file__).resolve().parents[1]
    / "0000737608_2026_08_ZL_04_742_03092026190352.xlsx"
)


def _xlsx_bytes(sheets):
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, rows, setup in sheets:
        sh = wb.active if first else wb.create_sheet(name)
        if first:
            sh.title = name
            first = False
        for row in rows:
            sh.append(list(row))
        if setup:
            setup(sh)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _one(rows, setup=None, name="Sheet1"):
    return _xlsx_bytes([(name, rows, setup)])


class TestExcelAliases(unittest.TestCase):
    def test_multiline_stock_headers(self):
        self.assertEqual(_xls_norm_header("PRODUCT DESCRIPTION"), "item")
        self.assertEqual(_xls_norm_header("OPENING\nSTOCK"), "op")
        self.assertEqual(_xls_norm_header("OPENING\nVALUE"), "opval")
        self.assertEqual(_xls_norm_header("RECEIVE\nQUANTITY"), "pur")
        self.assertEqual(_xls_norm_header("RECEIVE\nVALUE"), "purval")
        self.assertEqual(_xls_norm_header("ISSUE\nQUANTITY"), "sale")
        self.assertEqual(_xls_norm_header("ISSUE\nVALUE"), "sval")
        self.assertEqual(_xls_norm_header("CLOSING\nSTOCK"), "bal")
        self.assertEqual(_xls_norm_header("CLOSING\nVALUE"), "bval")
        self.assertNotEqual(_xls_norm_header("OPENING\nVALUE"), "op")
        self.assertNotEqual(_xls_norm_header("ISSUE\nVALUE"), "sale")
        self.assertNotEqual(_xls_norm_header("CLOSING\nVALUE"), "bal")

    def test_numeric_normalization(self):
        self.assertEqual(_to_float("1,250"), 1250.0)
        self.assertEqual(_to_float("₹ 1,250.50"), 1250.50)
        self.assertEqual(_to_float("(100)"), -100.0)
        self.assertEqual(_to_float("100.00"), 100.0)
        self.assertEqual(_to_float("100,00"), 100.0)
        self.assertEqual(_to_float("=SUM(H2:H20)"), 0.0)


class TestExcelLayouts(unittest.TestCase):
    def test_header_on_row_1(self):
        data = _one(
            [
                ["Item", "Op.", "Sale", "Bal."],
                ["LIV 52 TAB", 10, 2, 8],
            ]
        )
        result = extract_sales_statement(data, "row1.xlsx")
        self.assertEqual(result["totals"]["extra"]["header_row"], 0)
        self.assertEqual(result["line_items"][0]["sales_qty"], 2.0)

    def test_header_on_row_5_and_row_10(self):
        for blank_count, filename in ((4, "row5.xlsx"), (9, "row10.xlsx")):
            rows = [[None] for _ in range(blank_count)]
            rows.append(["Product Name", "Opening Qty", "Sales Qty", "Closing Qty"])
            rows.append(["ABANA TAB", 14, 5, 9])
            result = extract_sales_statement(_one(rows), filename)
            item = result["line_items"][0]
            self.assertEqual(item["product_name"], "ABANA TAB", filename)
            self.assertEqual(item["opening_qty"], 14.0, filename)
            self.assertEqual(item["sales_qty"], 5.0, filename)
            self.assertEqual(item["closing_qty"], 9.0, filename)
            self.assertEqual(result["totals"]["extra"]["header_row"], blank_count, filename)

    def test_hidden_sheet_is_not_used_when_visible_sheet_has_data(self):
        def hide(sh):
            sh.sheet_state = "hidden"

        data = _xlsx_bytes(
            [
                (
                    "Notes",
                    [
                        ["Item", "Op.", "Sale", "Bal."],
                        ["SHOULD NOT EXTRACT", 1, 1, 1],
                    ],
                    hide,
                ),
                (
                    "Stock",
                    [
                        ["Item Name", "Opening Stock", "Sales", "Closing Stock"],
                        ["LIV 52 TAB", 10, 2, 8],
                    ],
                    None,
                ),
            ]
        )
        result = extract_sales_statement(data, "hidden.xlsx")
        names = [i["product_name"] for i in result["line_items"]]
        self.assertEqual(names, ["LIV 52 TAB"])
        self.assertEqual(result["totals"]["extra"]["sheet"], "Stock")

    def test_multiline_and_reordered_headers(self):
        data = _one(
            [
                ["DRUGS POINT"],
                ["-HIMALAYA ZEAL STOCK & SALES STATEMENT 01-08-2026 - 31-08-2026"],
                [
                    "CLOSING\nSTOCK",
                    "CLOSING\nVALUE",
                    "ISSUE\nVALUE",
                    "ISSUE\nQUANTITY",
                    "RECEIVE\nQUANTITY",
                    "OPENING\nSTOCK",
                    "PRODUCT DESCRIPTION",
                ],
                [109, "17890.17", "3332.52", 18, 0, 127, "CONFIDO 60"],
                [0, 0, 0, 0, 0, 0, "AACTARIL SOAP 75 GM"],
                [881, "113844.7", "81318.58", 539, 240, 1180, "TOTAL"],
            ]
        )
        result = extract_sales_statement(data, "multiline.xlsx")
        self.assertEqual(result["stockist_name"], "DRUGS POINT")
        self.assertEqual(result["company_name"], "HIMALAYA ZEAL")
        self.assertEqual(result["period_from"], "2026-08-01")
        self.assertEqual(result["period_to"], "2026-08-31")
        names = [i["product_name"] for i in result["line_items"]]
        self.assertEqual(names, ["CONFIDO 60", "AACTARIL SOAP 75 GM"])
        confido = result["line_items"][0]
        self.assertEqual(confido["opening_qty"], 127.0)
        self.assertEqual(confido["receipts_qty"], 0.0)
        self.assertEqual(confido["sales_qty"], 18.0)
        self.assertEqual(confido["sales_value"], 3332.52)
        self.assertEqual(confido["closing_qty"], 109.0)
        self.assertEqual(confido["closing_value"], 17890.17)
        soap = result["line_items"][1]
        self.assertEqual(soap["opening_qty"], 0.0)
        self.assertEqual(soap["closing_value"], 0.0)
        self.assertEqual(result["totals"]["sales_value"], 81318.58)

    def test_multi_row_headers(self):
        data = _one(
            [
                ["Product", "Quantity", "Quantity", "Quantity"],
                [None, "Opening", "Sales", "Closing"],
                ["LIV 52 TAB", 10, 4, 6],
            ]
        )
        result = extract_sales_statement(data, "multirow.xlsx")
        item = result["line_items"][0]
        self.assertEqual(item["product_name"], "LIV 52 TAB")
        self.assertEqual(item["opening_qty"], 10.0)
        self.assertEqual(item["sales_qty"], 4.0)
        self.assertEqual(item["closing_qty"], 6.0)

    def test_merged_group_header(self):
        def merge(sh):
            sh.merge_cells("B1:D1")

        data = _one(
            [
                ["Product", "Quantity", None, None],
                [None, "Opening", "Sales", "Closing"],
                ["ABANA TAB", 9, 2, 7],
            ],
            setup=merge,
        )
        result = extract_sales_statement(data, "merged.xlsx")
        item = result["line_items"][0]
        self.assertEqual(item["opening_qty"], 9.0)
        self.assertEqual(item["sales_qty"], 2.0)
        self.assertEqual(item["closing_qty"], 7.0)

    def test_text_currency_dates_blanks_repeated_header_and_formula(self):
        data = _one(
            [
                ["Statement Date", "01/08/2026"],
                ["Item Code", "Item Name", "Opening", "Sales", "Sales Value", "Closing", "Closing Value"],
                ["000123", "LIV 52 TAB", "100", "1,250", "₹ 1,250.50", "80", 0],
                [],
                ["Item Code", "Item Name", "Opening", "Sales", "Sales Value", "Closing", "Closing Value"],
                ["000124", "CONFIDO 60", 1, "=SUM(H2:H20)", 10, 1, 5],
                ["TOTAL", "TOTAL", 101, 20, 100, 81, 5],
            ]
        )
        result = extract_sales_statement(data, "messy.xlsx")
        self.assertEqual(result["period_from"], "2026-08-01")
        names = [i["product_name"] for i in result["line_items"]]
        self.assertEqual(names, ["LIV 52 TAB", "CONFIDO 60"])
        liv = result["line_items"][0]
        self.assertEqual(liv["product_code"], "000123")
        self.assertEqual(liv["opening_qty"], 100.0)
        self.assertEqual(liv["sales_qty"], 1250.0)
        self.assertEqual(liv["sales_value"], 1250.50)
        self.assertEqual(liv["closing_value"], 0.0)
        confido = result["line_items"][1]
        self.assertNotEqual(confido["sales_qty"], 2.0)
        self.assertTrue(confido["extra"].get("unresolved_formula"))
        self.assertIsNone(result.get("stockist_name"))

    def test_missing_optional_columns_keep_schema(self):
        data = _one(
            [
                ["Product", "Sales Qty"],
                ["ABANA TAB", 3],
            ]
        )
        result = extract_sales_statement(data, "partial.xlsx")
        item = result["line_items"][0]
        self.assertEqual(item["sales_qty"], 3.0)
        self.assertEqual(item["opening_qty"], 0.0)
        self.assertIsNone(result["period_from"])
        mapping = result["totals"]["extra"]["column_mapping"]
        self.assertEqual(mapping["opening_qty"]["confidence"], 0.0)
        self.assertGreater(mapping["sales_qty"]["confidence"], 0)

    def test_multiple_sheets_keep_visible_product_rows(self):
        data = _xlsx_bytes(
            [
                ("Cover", [["Notes only"]], None),
                (
                    "Summary",
                    [
                        ["Item", "Opening", "Sales", "Closing"],
                        ["LIV 52 TAB", 1, 1, 1],
                    ],
                    None,
                ),
                (
                    "Detail",
                    [
                        ["Item Name", "Opening Qty", "Sales Qty", "Closing Qty"],
                        ["LIV 52 TAB", 10, 2, 8],
                        ["ABANA TAB", 4, 1, 3],
                    ],
                    None,
                ),
            ]
        )
        result = extract_sales_statement(data, "sheets.xlsx")
        names = [i["product_name"] for i in result["line_items"]]
        self.assertIn("ABANA TAB", names)
        self.assertIn("LIV 52 TAB", names)
        liv = [i for i in result["line_items"] if i["product_name"] == "LIV 52 TAB"]
        self.assertEqual(len(liv), 1)
        self.assertEqual(liv[0]["opening_qty"], 10.0)
        self.assertEqual(result["totals"]["extra"]["sheet"], "Detail")


class TestProvidedXls(unittest.TestCase):
    def test_himalaya_drugs_point_xls(self):
        if not SAMPLE_XLS.is_file():
            self.skipTest(f"missing {SAMPLE_XLS}")
        result = extract_sales_statement(SAMPLE_XLS.read_bytes(), SAMPLE_XLS.name)
        extra = result["totals"]["extra"]
        self.assertEqual(extra.get("sheet"), "Stock & Sales Statement")
        self.assertEqual(extra.get("header_row"), 6)
        self.assertEqual(result["stockist_name"], "DRUGS POINT")
        self.assertEqual(len(result["line_items"]), 23)
        by_name = {i["product_name"]: i for i in result["line_items"]}
        confido = by_name["CONFIDO 60"]
        self.assertEqual(confido["opening_qty"], 127.0)
        self.assertEqual(confido["receipts_qty"], 0.0)
        self.assertEqual(confido["sales_qty"], 18.0)
        self.assertAlmostEqual(confido["sales_value"], 3332.52)
        self.assertEqual(confido["closing_qty"], 109.0)
        self.assertAlmostEqual(confido["closing_value"], 17890.17)
        self.assertTrue(confido["extra"].get("stock_identity_ok"))
        liv = by_name["LIV 52 SYP 200 ML"]
        self.assertEqual(liv["opening_qty"], 127.0)
        self.assertEqual(liv["receipts_qty"], 140.0)
        self.assertEqual(liv["sales_qty"], 171.0)
        self.assertEqual(liv["closing_qty"], 96.0)
        soap = by_name["AACTARIL SOAP 75 GM"]
        self.assertEqual(soap["opening_qty"], 0.0)
        self.assertEqual(soap["sales_qty"], 0.0)
        self.assertEqual(soap["closing_qty"], 0.0)
        self.assertEqual(soap["closing_value"], 0.0)
        self.assertNotIn("TOTAL", by_name)
        self.assertAlmostEqual(result["totals"]["sales_value"], 81318.58)
        self.assertAlmostEqual(result["totals"]["closing_value"], 113844.7)
        mapping = extra["column_mapping"]
        self.assertEqual(mapping["product"]["header"], "PRODUCT DESCRIPTION")
        self.assertEqual(mapping["sales_qty"]["header"], "ISSUE QUANTITY")
        self.assertEqual(mapping["closing_value"]["header"], "CLOSING VALUE")


class TestKnownXlsx(unittest.TestCase):
    def test_previously_working_mat_name_workbook(self):
        if not KNOWN_XLSX.is_file():
            self.skipTest(f"missing {KNOWN_XLSX}")
        result = extract_sales_statement(KNOWN_XLSX.read_bytes(), KNOWN_XLSX.name)
        self.assertEqual(result["stockist_name"], "RICHA PHARMA")
        coded = {i["product_code"]: i for i in result["line_items"] if i.get("product_code")}
        kap = coded["7002675"]
        self.assertEqual(kap["product_name"], "KAPIKACHHU TABS 60 S INDIA")
        self.assertEqual(kap["receipts_qty"], 60.0)
        self.assertEqual(kap["sales_qty"], 30.0)
        self.assertEqual(kap["closing_qty"], 30.0)


if __name__ == "__main__":
    unittest.main()
