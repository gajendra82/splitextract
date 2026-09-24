"""Excel stock-statement period date extraction / normalization."""

from __future__ import annotations

import io
import unittest
from datetime import date, datetime
from pathlib import Path

from services.sales_statement_extractor import (
    _apply_excel_header_period,
    _normalize_date,
    _normalize_excel_date,
    _drop_trailing_statement_total_item,
    _finalize_zandra_stock_sale,
    _looks_like_zandra_stock_sale_text,
    _parse_ps_pharma_statement,
    _parse_rate_qty_value_statement,
    _parse_semantic_text_table,
    _pdf_pages_are_image_only,
    _statement_numbers_are_blank,
    _statement_nonzero_rows,
    _to_float,
    _xls_best_header,
    _xls_header_colmap,
    _xls_is_stock_header,
    _xls_norm_header,
    empty_result,
    extract_sales_statement,
)


SAMPLE_XLSX = (
    Path(__file__).resolve().parents[1]
    / "0000737569_2026_08_ZL_13_285_01092026073322.xlsx"
)


class TestNormalizeExcelDate(unittest.TestCase):
    def test_datetime_cell(self):
        self.assertEqual(
            _normalize_excel_date(datetime(2026, 8, 1, 13, 45, 0)),
            "2026-08-01",
        )

    def test_date_cell(self):
        self.assertEqual(_normalize_excel_date(date(2026, 8, 29)), "2026-08-29")

    def test_excel_serial_with_date_format(self):
        # 2026-08-01 = 46235 in the 1899-12-30 Excel epoch
        self.assertEqual(
            _normalize_excel_date(46235, number_format="DD/MM/YYYY"),
            "2026-08-01",
        )
        self.assertEqual(
            _normalize_excel_date(46235.0, number_format="YYYY-MM-DD"),
            "2026-08-01",
        )

    def test_dd_mm_yyyy_string(self):
        self.assertEqual(_normalize_excel_date("01/08/2026"), "2026-08-01")
        self.assertEqual(_normalize_date("01/08/2026"), "2026-08-01")

    def test_dd_mm_yyyy_dashed_string(self):
        self.assertEqual(_normalize_excel_date("01-08-2026"), "2026-08-01")

    def test_yyyy_mm_dd_string(self):
        self.assertEqual(_normalize_excel_date("2026-08-01"), "2026-08-01")
        self.assertEqual(_normalize_date("2026-08-01"), "2026-08-01")
        self.assertEqual(_normalize_date("2026-08-01 00:00:00"), "2026-08-01")

    def test_yyyy_mm_dd_slashed(self):
        self.assertEqual(_normalize_date("2026/08/01"), "2026-08-01")

    def test_empty_date(self):
        self.assertIsNone(_normalize_excel_date(None))
        self.assertIsNone(_normalize_excel_date(""))
        self.assertIsNone(_normalize_excel_date("   "))
        self.assertIsNone(_normalize_date(None))
        self.assertIsNone(_normalize_date(""))

    def test_invalid_date(self):
        self.assertIsNone(_normalize_excel_date("not-a-date"))
        self.assertIsNone(_normalize_excel_date("32/13/2026"))
        self.assertIsNone(_normalize_date("Report Date"))

    def test_numeric_quantity_is_not_a_date(self):
        self.assertIsNone(_normalize_excel_date(8))
        self.assertIsNone(_normalize_excel_date(74, number_format="General"))
        self.assertIsNone(_normalize_excel_date(269))

    def test_sales_value_is_not_a_date(self):
        self.assertIsNone(_normalize_excel_date(11983.26, number_format="0.00"))
        self.assertIsNone(_normalize_excel_date(11983.26, number_format="#,##0.00"))
        self.assertIsNone(_normalize_excel_date(11983.26))

    def test_pandas_timestamp(self):
        try:
            import pandas as pd
        except ImportError:
            self.skipTest("pandas not installed")
        self.assertEqual(
            _normalize_excel_date(pd.Timestamp("2026-08-01 07:33:22")),
            "2026-08-01",
        )
        self.assertIsNone(_normalize_excel_date(pd.NaT))


class TestExcelHeaderPeriod(unittest.TestCase):
    def test_from_to_slash_strings_in_title(self):
        result = empty_result("stmt.xlsx", "xlsx")
        _apply_excel_header_period(
            result,
            [["Sales And Stock Register From 01/08/2026 To 29/08/2026"]],
        )
        self.assertEqual(result["period_from"], "2026-08-01")
        self.assertEqual(result["period_to"], "2026-08-29")

    def test_from_and_to_on_separate_rows(self):
        result = empty_result("stmt.xlsx", "xlsx")
        _apply_excel_header_period(
            result,
            [
                ["From : 01/08/2026", "-", "-"],
                ["To : 29/08/2026", "-", "-"],
            ],
        )
        self.assertEqual(result["period_from"], "2026-08-01")
        self.assertEqual(result["period_to"], "2026-08-29")

    def test_datetime_cell_beside_from_to_labels(self):
        result = empty_result("stmt.xlsx", "xlsx")
        _apply_excel_header_period(
            result,
            [
                ["From :", datetime(2026, 8, 1)],
                ["To :", datetime(2026, 8, 29)],
            ],
        )
        self.assertEqual(result["period_from"], "2026-08-01")
        self.assertEqual(result["period_to"], "2026-08-29")

    def test_does_not_take_first_date_from_item_rows(self):
        result = empty_result("stmt.xlsx", "xlsx")
        rows = [
            ["SIND AGENCIES"],
            ["Item Name", "Sale"],
        ] + [["PROD %s" % i, 10] for i in range(25)]
        _apply_excel_header_period(result, rows)
        self.assertIsNone(result["period_from"])
        self.assertIsNone(result["period_to"])

    def test_does_not_overwrite_existing_period(self):
        result = empty_result("stmt.xlsx", "xlsx")
        result["period_from"] = "2026-07-01"
        result["period_to"] = "2026-07-31"
        _apply_excel_header_period(
            result, [["From : 01/08/2026 To 29/08/2026"]]
        )
        self.assertEqual(result["period_from"], "2026-07-01")
        self.assertEqual(result["period_to"], "2026-07-31")


class TestSindAgenciesXlsxRegression(unittest.TestCase):
    def test_sample_xlsx_period_from_content_not_filename(self):
        self.assertTrue(SAMPLE_XLSX.is_file(), f"missing fixture {SAMPLE_XLSX}")
        data = SAMPLE_XLSX.read_bytes()
        result = extract_sales_statement(
            data, "0000737569_2026_08_ZL_13_285_01092026073322.xlsx"
        )
        self.assertEqual(result["period_from"], "2026-08-01")
        self.assertEqual(result["period_to"], "2026-08-29")
        # Filename stamp is 01/09/2026 07:33:22 — must not become the period.
        self.assertNotEqual(result["period_from"], "2026-09-01")
        self.assertNotEqual(result["period_to"], "2026-09-01")
        self.assertEqual(result["stockist_name"], "SIND AGENCIES")
        self.assertGreater(len(result.get("line_items") or []), 0)
        names = [i.get("product_name") for i in result["line_items"]]
        self.assertNotIn("SIND AGENCIES", names)
        self.assertFalse(any(n and str(n).startswith("VALUE OF") for n in names))
        soap = next(i for i in result["line_items"] if i["product_name"] == "*AACTARIL SOAP")
        self.assertEqual(soap["packing"], "75GM")
        self.assertEqual(soap["opening_qty"], 20.0)
        self.assertEqual(soap["receipts_qty"], 0.0)
        self.assertEqual(soap["sales_qty"], 1.0)
        self.assertEqual(soap["closing_qty"], 19.0)
        abana = next(i for i in result["line_items"] if i["product_name"] == "ABANA TAB")
        self.assertEqual(abana["packing"], "60 S")
        self.assertEqual(abana["opening_qty"], 14.0)
        self.assertEqual(abana["sales_qty"], 5.0)
        self.assertEqual(abana["closing_qty"], 9.0)
        extra = (result.get("totals") or {}).get("extra") or {}
        printed_sale = extra.get("value_of_sale") or extra.get("rejected_sales_value")
        self.assertAlmostEqual(float(printed_sale), 53076.66)
        self.assertAlmostEqual(result["totals"]["closing_value"], 126298.16)
        self.assertFalse(
            any(n and str(n).upper().startswith("TOTAL") for n in names)
        )


class TestExcelDatetimeWorkbook(unittest.TestCase):
    def test_openpyxl_datetime_cells_in_header(self):
        from openpyxl import Workbook

        wb = Workbook()
        sh = wb.active
        sh["A1"] = "TEST AGENCIES"
        sh["A2"] = "From :"
        sh["B2"] = datetime(2026, 8, 1)
        sh["B2"].number_format = "DD/MM/YYYY"
        sh["A3"] = "To :"
        sh["B3"] = datetime(2026, 8, 29)
        sh["B3"].number_format = "DD/MM/YYYY"
        sh["A4"] = "Item"
        sh["B4"] = "Op."
        sh["C4"] = "Sale"
        sh["D4"] = "Bal."
        sh["A5"] = "LIV 52 TAB"
        sh["B5"] = 10
        sh["C5"] = 2
        sh["D5"] = 8
        buf = io.BytesIO()
        wb.save(buf)
        result = extract_sales_statement(buf.getvalue(), "datetime-header.xlsx")
        self.assertEqual(result["period_from"], "2026-08-01")
        self.assertEqual(result["period_to"], "2026-08-29")
        liv = next(
            i for i in result["line_items"] if "LIV 52" in (i.get("product_name") or "")
        )
        self.assertEqual(liv["opening_qty"], 10.0)
        self.assertEqual(liv["sales_qty"], 2.0)
        self.assertEqual(liv["closing_qty"], 8.0)


MAT_NAME_XLSX = (
    Path(__file__).resolve().parents[1]
    / "0000737608_2026_08_ZL_04_742_03092026190352.xlsx"
)


class TestMatNameExcelHeaders(unittest.TestCase):
    def test_mat_name_aliases(self):
        for label in (
            "Mat Name",
            "MAT NAME",
            "mat name",
            "Mat  Name",
            "Mat\nName",
            "MAT\nNAME",
            "Mat\u00a0Name",
        ):
            self.assertEqual(_xls_norm_header(label), "item", label)

    def test_mat_code_aliases(self):
        self.assertEqual(_xls_norm_header("Mat Code"), "product_code")
        self.assertEqual(_xls_norm_header("MAT CODE"), "product_code")

    def test_mat_is_not_confused_with_name(self):
        self.assertNotEqual(_xls_norm_header("Mat Code"), "item")
        self.assertNotEqual(_xls_norm_header("Customer Name"), "item")

    def test_header_row_detected(self):
        row = [
            "Year",
            "Month",
            "Div",
            "Cust Code",
            "Customer Name",
            "Mat Code",
            "Mat Name",
            "Secondary Rate",
            "Ob (Qty)",
            "Primary (Qty)",
            "Cb (Qty)",
            "Secondary Qty Total",
        ]
        colmap = _xls_header_colmap(row)
        self.assertTrue(_xls_is_stock_header(colmap))
        self.assertEqual(colmap["item"], 6)
        self.assertEqual(colmap["product_code"], 5)
        self.assertEqual(colmap["customer_name"], 4)
        self.assertEqual(colmap["cust_code"], 3)
        self.assertEqual(colmap["op"], 8)
        self.assertEqual(colmap["pur"], 9)
        self.assertEqual(colmap["bal"], 10)
        self.assertEqual(colmap["sale"], 11)

    def test_sample_xlsx_mat_name_and_code(self):
        self.assertTrue(MAT_NAME_XLSX.is_file(), f"missing fixture {MAT_NAME_XLSX}")
        result = extract_sales_statement(MAT_NAME_XLSX.read_bytes(), MAT_NAME_XLSX.name)
        self.assertEqual(result["stockist_name"], "RICHA PHARMA")
        self.assertEqual(result["totals"]["extra"].get("stockist_code"), "0000737608")
        self.assertEqual(result.get("stockist_code"), "0000737608")
        self.assertEqual(result["totals"]["extra"].get("division"), "ZL")
        self.assertEqual(result.get("division"), "ZL")
        # Year/Month=2026/8 is not a calendar date. Do not invent 2026-08-01.
        self.assertIsNone(result["period_from"])
        self.assertIsNone(result["period_to"])
        self.assertIsNone(result.get("statement_date"))
        self.assertEqual(result["totals"]["extra"].get("year"), "2026")
        self.assertEqual(result["totals"]["extra"].get("month"), "8")
        coded = [
            i
            for i in result["line_items"]
            if i.get("product_code") and i.get("product_name")
        ]
        self.assertGreaterEqual(len(coded), 5)
        by_code = {i["product_code"]: i for i in coded}
        self.assertEqual(by_code["7002675"]["product_name"], "KAPIKACHHU TABS 60 S INDIA")
        self.assertEqual(
            by_code["7006064"]["product_name"], "LIV.52 SUGAR FREE SYRUP 200ML IND"
        )
        self.assertEqual(by_code["7006097"]["product_name"], "LIV.52 TABS (FC) 100 S")
        self.assertEqual(by_code["7002319"]["product_name"], "AACTARIL SOAP 75G INDIA")
        self.assertEqual(by_code["7000004"]["product_name"], "ABANA TABS 60 s")
        kap = by_code["7002675"]
        self.assertEqual(kap["opening_qty"], 0.0)
        self.assertEqual(kap["receipts_qty"], 60.0)
        self.assertEqual(kap["closing_qty"], 30.0)
        self.assertEqual(kap["sales_qty"], 30.0)
        self.assertNotEqual(kap["product_name"], "RICHA PHARMA")
        self.assertNotEqual(kap["product_code"], "KAPIKACHHU TABS 60 S INDIA")
        self.assertFalse(
            any(i.get("product_name") == "RICHA PHARMA" for i in result["line_items"])
        )
        self.assertAlmostEqual(kap["extra"].get("unit_rate"), 203.66)
        meta = (result.get("totals") or {}).get("extra") or {}
        self.assertEqual(meta.get("product_column_header"), "Mat Name")
        self.assertIn("stock_validation", meta)


class TestSemanticHeaderAliases(unittest.TestCase):
    def test_product_name_aliases(self):
        for label in (
            "Product Name",
            "Product",
            "Product Description",
            "Product Desc",
            "Item Name",
            "Item",
            "Material Name",
            "Mat Name",
            "Material",
            "Medicine Name",
            "Drug Name",
            "SKU Description",
            "Description",
        ):
            self.assertEqual(_xls_norm_header(label), "item", label)

    def test_qty_and_party_aliases(self):
        self.assertEqual(_xls_norm_header("Ob Qty"), "op")
        self.assertEqual(_xls_norm_header("Ob"), "op")
        self.assertEqual(_xls_norm_header("Primary Qty"), "pur")
        self.assertEqual(_xls_norm_header("Cb Qty"), "bal")
        self.assertEqual(_xls_norm_header("Secondary Qty Total"), "sale")
        self.assertEqual(_xls_norm_header("Cust Code"), "cust_code")
        self.assertEqual(_xls_norm_header("Party Name"), "customer_name")
        self.assertEqual(_xls_norm_header("PTR"), "rate")
        self.assertEqual(_xls_norm_header("Div"), "div")

    def test_value_is_not_quantity(self):
        self.assertEqual(_xls_norm_header("Secondary Value"), "sval")
        self.assertNotEqual(_xls_norm_header("Secondary Value"), "sale")
        self.assertNotEqual(_xls_norm_header("Secondary Rate"), "sale")

    def test_header_scoring_skips_title_rows(self):
        rows = [
            ["HIMALAYA STOCK STATEMENT"],
            ["Generated on 01/08/2026"],
            [],
            [
                "Mat Code",
                "Mat Name",
                "Secondary Rate",
                "Ob (Qty)",
                "Cb (Qty)",
            ],
            ["7002675", "KAPIKACHHU TABS 60 S INDIA", 203.66, 0, 30],
        ]
        idx, colmap, score = _xls_best_header(rows)
        self.assertEqual(idx, 3)
        self.assertEqual(colmap["item"], 1)
        self.assertGreater(score, 0)


class TestExcelLayoutVariations(unittest.TestCase):
    def _xlsx(self, rows, sheets=None):
        from openpyxl import Workbook

        wb = Workbook()
        if sheets is None:
            sheets = [("Sheet1", rows)]
        first = True
        for name, data in sheets:
            sh = wb.active if first else wb.create_sheet(name)
            if first:
                sh.title = name
                first = False
            for r in data:
                sh.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_column_order_does_not_matter(self):
        data = self._xlsx(
            [
                ["Cb (Qty)", "Mat Name", "Ob (Qty)", "Secondary Qty Total", "Primary (Qty)"],
                [8, "LIV 52 TAB", 10, 2, 0],
            ]
        )
        result = extract_sales_statement(data, "reorder.xlsx")
        liv = result["line_items"][0]
        self.assertEqual(liv["product_name"], "LIV 52 TAB")
        self.assertEqual(liv["opening_qty"], 10.0)
        self.assertEqual(liv["receipts_qty"], 0.0)
        self.assertEqual(liv["sales_qty"], 2.0)
        self.assertEqual(liv["closing_qty"], 8.0)

    def test_title_and_blank_rows_before_header(self):
        data = self._xlsx(
            [
                ["ACME PHARMA STOCK REPORT"],
                [None, None],
                [],
                ["Item Name", "Op. Stock", "Purch", "Sale", "Cl Stock"],
                ["ABANA TAB", 14, 0, 5, 9],
            ]
        )
        result = extract_sales_statement(data, "titles.xlsx")
        self.assertEqual(result["line_items"][0]["product_name"], "ABANA TAB")
        self.assertEqual(result["line_items"][0]["opening_qty"], 14.0)
        self.assertEqual(result["line_items"][0]["closing_qty"], 9.0)
        self.assertNotIn("ACME PHARMA STOCK REPORT", [i["product_name"] for i in result["line_items"]])

    def test_multiple_sheets_combines_valid_tables(self):
        data = self._xlsx(
            None,
            sheets=[
                ("Cover", [["Notes"], ["Ignore this sheet"]]),
                (
                    "Stock",
                    [
                        ["Item", "Op.", "Sale", "Bal."],
                        ["LIV 52 TAB", 10, 2, 8],
                    ],
                ),
                (
                    "More",
                    [
                        ["Item Name", "Op. Stock", "Sale", "Cl Stock"],
                        ["ABANA TAB", 4, 1, 3],
                    ],
                ),
            ],
        )
        result = extract_sales_statement(data, "multisheet.xlsx")
        names = [i["product_name"] for i in result["line_items"]]
        self.assertIn("LIV 52 TAB", names)
        self.assertIn("ABANA TAB", names)
        self.assertEqual(len(result["line_items"]), 2)

    def test_missing_date_stays_null(self):
        data = self._xlsx(
            [
                ["Item", "Op.", "Sale", "Bal."],
                ["LIV 52 TAB", 10, 2, 8],
            ]
        )
        result = extract_sales_statement(data, "nodate.xlsx")
        self.assertIsNone(result["period_from"])
        self.assertIsNone(result["period_to"])
        self.assertIsNone(result.get("statement_date"))

    def test_numeric_parsing(self):
        self.assertEqual(_to_float("1,250"), 1250.0)
        self.assertEqual(_to_float("₹1,250.00"), 1250.0)
        self.assertEqual(_to_float("1 250"), 1250.0)
        self.assertEqual(_to_float("(125)"), -125.0)
        self.assertEqual(_to_float("--"), 0.0)
        self.assertEqual(_to_float("N/A"), 0.0)

    def test_semantic_pdf_text_table(self):
        text = (
            "RICHA PHARMA\n"
            "Cust Code\tCustomer Name\tMat Code\tMat Name\tOb (Qty)\tPrimary (Qty)\t"
            "Cb (Qty)\tSecondary Qty Total\n"
            "0000737608\tRICHA PHARMA\t7002675\tKAPIKACHHU TABS 60 S INDIA\t0\t60\t30\t30\n"
        )
        result = _parse_semantic_text_table(text, "stmt.pdf", "pdf")
        self.assertIsNotNone(result)
        self.assertEqual(result["line_items"][0]["product_name"], "KAPIKACHHU TABS 60 S INDIA")
        self.assertEqual(result["line_items"][0]["product_code"], "7002675")
        self.assertEqual(result["line_items"][0]["receipts_qty"], 60.0)
        self.assertEqual(result["line_items"][0]["sales_qty"], 30.0)


    def test_op_purch_sale_others_closing_stay_on_source_columns(self):
        data = self._xlsx(
            [
                ["Product Name", "Pack", "Op. Stock", "Purch", "Sale", "#Others", "Cl Stock"],
                ["RUMALAYA (LINIMENT)", "60ML", 8, -2, 4, 1, 3],
                ["LIV.52 SYRUP", "100ML", 0, 0, 0, 0, 4],
            ]
        )
        result = extract_sales_statement(data, "op-purch-sale.xlsx")
        rumalaya, liv = result["line_items"]
        self.assertEqual(rumalaya["product_name"], "RUMALAYA (LINIMENT)")
        self.assertEqual(rumalaya["packing"], "60ML")
        self.assertEqual(rumalaya["opening_qty"], 8.0)
        self.assertEqual(rumalaya["receipts_qty"], -2.0)
        self.assertEqual(rumalaya["sales_qty"], 4.0)
        self.assertEqual(rumalaya["closing_qty"], 3.0)
        self.assertEqual(rumalaya["extra"]["others_qty"], 1.0)
        self.assertEqual(liv["opening_qty"], 0.0)
        self.assertEqual(liv["sales_qty"], 0.0)
        self.assertEqual(liv["closing_qty"], 4.0)
        self.assertNotEqual(liv["sales_qty"], liv["closing_qty"])


class TestScannedPdfRouting(unittest.TestCase):
    def test_image_only_pages_detected(self):
        self.assertTrue(
            _pdf_pages_are_image_only(
                [{"words": []}, {"words": None}]
            )
        )
        self.assertFalse(
            _pdf_pages_are_image_only(
                [{"words": [(0, 0, 1, 1, "Item")]}]
            )
        )

    def test_gemini_zero_qty_rows_are_blank(self):
        result = empty_result("scan.pdf", "pdf")
        result["line_items"] = [
            {
                "product_name": "EVECARE CAPSULE",
                "opening_qty": 0,
                "receipts_qty": 0,
                "sales_qty": 0,
                "sales_value": 0,
                "closing_qty": 0,
                "closing_value": 0,
            }
            for _ in range(10)
        ]
        result["line_items"][0]["opening_qty"] = 130
        self.assertEqual(_statement_nonzero_rows(result), 1)
        self.assertTrue(_statement_numbers_are_blank(result))

    def test_trailing_footer_totals_dropped(self):
        items = [
            {"product_name": "ARJUNA TABLET", "opening_qty": 49, "receipts_qty": 20}
            for _ in range(10)
        ]
        items.append(
            {
                "product_name": "V-GEL CREAM",
                "opening_qty": 0,
                "receipts_qty": 2961,
                "sales_qty": 658,
                "sales_value": 78083,
            }
        )
        trimmed = _drop_trailing_statement_total_item(items)
        self.assertEqual(len(trimmed), 10)
        self.assertEqual(trimmed[-1]["product_name"], "ARJUNA TABLET")

    def test_zandra_header_detected(self):
        self.assertTrue(
            _looks_like_zandra_stock_sale_text(
                "BINAL PHARMA\nStock and Sale Statement\nItem Cd Item Name Op Stk"
            )
        )

    def test_zandra_skips_division_row(self):
        result = empty_result("scan.png", "png")
        result["line_items"] = [
            {"product_name": "HIMALAYA ZANDRA DIVISION-ZAN", "opening_qty": 0},
            {
                "product_name": "ARJUNA TABLET",
                "opening_qty": 49,
                "sales_qty": 2,
                "closing_qty": 47,
                "sales_value": 495,
                "closing_value": 9333,
                "extra": {},
            },
        ]
        result = _finalize_zandra_stock_sale(result)
        names = [i["product_name"] for i in result["line_items"]]
        self.assertEqual(names, ["ARJUNA TABLET"])
        self.assertEqual(result["line_items"][0]["closing_qty"], 47)
        self.assertEqual(result["line_items"][0]["closing_value"], 9333)

class TestSecondarySalesQtyMapping(unittest.TestCase):
    def test_sale_zero_is_not_replaced_by_closing_stock(self):
        from app import _sales_line_to_invoice_item

        item = _sales_line_to_invoice_item({
            "product_name": "RUMALAYA (LINIMENT)",
            "packing": "60ML",
            "opening_qty": 0.0,
            "receipts_qty": 0.0,
            "sales_qty": 0.0,
            "closing_qty": 4.0,
            "extra": {"others_qty": 4.0},
        })
        self.assertEqual(item["quantity"], "0.0")
        fields = item["additional_fields"]
        self.assertEqual(fields["opening_qty"], 0.0)
        self.assertEqual(fields["receipts_qty"], 0.0)
        self.assertEqual(fields["sales_qty"], 0.0)
        self.assertEqual(fields["others_qty"], 4.0)
        self.assertEqual(fields["closing_qty"], 4.0)
        self.assertEqual(fields["packing"], "60ML")

    def test_source_columns_stay_on_their_own_fields(self):
        from app import _sales_line_to_invoice_item

        item = _sales_line_to_invoice_item({
            "product_name": "RUMALAYA (LINIMENT)",
            "packing": "60ML",
            "opening_qty": 8.0,
            "receipts_qty": -2.0,
            "sales_qty": 4.0,
            "closing_qty": 3.0,
            "extra": {"others_qty": 1.0},
        })
        self.assertEqual(item["quantity"], "4.0")
        fields = item["additional_fields"]
        self.assertEqual(fields["opening_qty"], 8.0)
        self.assertEqual(fields["receipts_qty"], -2.0)
        self.assertEqual(fields["sales_qty"], 4.0)
        self.assertEqual(fields["others_qty"], 1.0)
        self.assertEqual(fields["closing_qty"], 3.0)


class TestPsPharmaDashCells(unittest.TestCase):
    def test_dash_cells_stay_zero_and_printed_closing_is_kept(self):
        text = """LAXMI MEDICAL DISTRIBUTOR
12-6,FIRST FLOOR
MAIN ROAD
MADHIRA - 507203
STOCK & SALES ANALYSIS 01/08/2026 - 31/08/2026
ITEM DESCRIPTION OPENING RECEIPT ISSUE CLOSING
THE HIMALAYA DRUG CO
RUMALAYA (LINIMENT)        60ML              7         -         4         3
BONNISAN LIQ  200`ML
140 - 84 56
GASEX TAB 100`S 58 - 4 54
TOTAL 286585 24676 253001 74263
"""
        result = _parse_ps_pharma_statement(text, "laxmi.pdf")
        self.assertIsNotNone(result)
        by_name = {
            (item["product_name"], item.get("packing")): item
            for item in result["line_items"]
        }
        rumalaya = by_name[("RUMALAYA (LINIMENT)", "60ML")]
        self.assertEqual(rumalaya["opening_qty"], 7.0)
        self.assertEqual(rumalaya["receipts_qty"], 0.0)
        self.assertEqual(rumalaya["sales_qty"], 4.0)
        self.assertEqual(rumalaya["closing_qty"], 3.0)
        wrapped = by_name[("BONNISAN LIQ", "200`ML")]
        self.assertEqual(
            (wrapped["opening_qty"], wrapped["receipts_qty"], wrapped["sales_qty"], wrapped["closing_qty"]),
            (140.0, 0.0, 84.0, 56.0),
        )
        gasex = by_name[("GASEX TAB", "100`S")]
        self.assertEqual(gasex["sales_qty"], 4.0)
        self.assertEqual(gasex["opening_qty"], 58.0)
        self.assertEqual(result["totals"]["opening_qty"], 286585.0)
        self.assertEqual(result["totals"]["closing_qty"], 74263.0)


class TestRateQtyValueColumns(unittest.TestCase):
    def test_blank_receipt_and_issue_stay_in_their_columns(self):
        def word(x0, x1, y, text):
            return (x0, y, x1, y + 8, text)

        header = [
            word(20, 55, 0, "STOCK"),
            word(58, 66, 0, "&"),
            word(70, 110, 0, "SALES"),
            word(20, 42, 10, "ITEM"),
            word(48, 107, 10, "DESCRIPTION"),
            word(172, 194, 10, "RATE"),
            word(242, 280, 10, "OPENING"),
            word(356, 393, 10, "RECEIPT"),
            word(464, 491, 10, "ISSUE"),
            word(572, 609, 10, "CLOSING"),
            word(647, 669, 10, "DUMP"),
        ]
        sub = [
            word(221, 242, 20, "QTY."),
            word(280, 307, 20, "VALUE"),
            word(334, 356, 20, "QTY."),
            word(393, 420, 20, "VALUE"),
            word(442, 464, 20, "QTY."),
            word(501, 528, 20, "VALUE"),
            word(550, 572, 20, "QTY."),
            word(609, 636, 20, "VALUE"),
            word(647, 669, 20, "QTY."),
        ]
        syrup = [
            word(21, 64, 40, "GERIFORT"),
            word(69, 86, 40, "SYP"),
            word(91, 113, 40, "200M"),
            word(118, 145, 40, "200ML"),
            word(161, 194, 40, "152.47"),
            word(226, 237, 40, "32"),
            word(264, 302, 40, "4879.06"),
            word(345, 350, 40, "-"),
            word(393, 415, 40, "0.00"),
            word(453, 458, 40, "-"),
            word(501, 523, 40, "0.00"),
            word(555, 566, 40, "32"),
            word(593, 631, 40, "4879.06"),
        ]
        tab = [
            word(21, 64, 52, "GERIFORT"),
            word(69, 86, 52, "TAB"),
            word(91, 107, 52, "100"),
            word(118, 150, 52, "100TAB"),
            word(161, 194, 52, "108.18"),
            word(226, 237, 52, "38"),
            word(264, 302, 52, "4110.71"),
            word(345, 350, 52, "-"),
            word(393, 415, 52, "0.00"),
            word(453, 458, 52, "-"),
            word(501, 523, 52, "0.00"),
            word(555, 566, 52, "38"),
            word(593, 631, 52, "4110.71"),
        ]
        nxt = [
            word(21, 69, 64, "HIMPLASIA"),
            word(75, 91, 64, "TAB"),
            word(118, 140, 64, "1*30"),
            word(161, 194, 64, "347.52"),
            word(226, 237, 64, "-5"),
            word(258, 302, 64, "-1737.59"),
            word(345, 350, 64, "-"),
            word(393, 415, 64, "0.00"),
            word(453, 458, 64, "8"),
            word(485, 523, 64, "2775.74"),
            word(550, 566, 64, "-13"),
            word(588, 631, 64, "-4517.74"),
        ]
        text = "KAPEESH MEDICAL STORE\nSTOCK & SALES ANALYSIS 01-08-2026 - 31-08-2026\n"
        pages = [{
            "text": text,
            "words": header + sub + syrup + tab + nxt,
        }]
        result = _parse_rate_qty_value_statement(pages, "kapeesh.pdf")
        self.assertIsNotNone(result)
        by_pack = {item.get("packing"): item for item in result["line_items"]}
        tab_row = by_pack["100TAB"]
        self.assertEqual(tab_row["product_name"], "GERIFORT TAB 100")
        self.assertEqual(tab_row["opening_qty"], 38.0)
        self.assertEqual(tab_row["receipts_qty"], 0.0)
        self.assertEqual(tab_row["sales_qty"], 0.0)
        self.assertEqual(tab_row["closing_qty"], 38.0)
        syrup_row = by_pack["200ML"]
        self.assertEqual(
            (syrup_row["opening_qty"], syrup_row["closing_qty"]),
            (32.0, 32.0),
        )
        neighbor = by_pack["1*30"]
        self.assertEqual(neighbor["opening_qty"], -5.0)
        self.assertEqual(neighbor["sales_qty"], 8.0)
        self.assertEqual(neighbor["closing_qty"], -13.0)


if __name__ == "__main__":
    unittest.main()
