"""Phase 4a STOCK_WORD_LINUX — LibreOffice .doc convert + image-only Word (offline)."""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from services.sales_statement_extractor import (
    empty_result,
    _convert_doc_to_docx_bytes_linux,
    _docx_extract_media_images,
    _docx_try_vision_images,
    _parse_word,
    _stock_word_linux_enabled,
)
from services.extraction_quality import assess_extraction_quality
from tests.gemini_offline import ensure_gemini_offline_guard, reset_gemini_test_call_count


def _soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def _make_table_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("ALPHA MEDICAL STORES")
    table = doc.add_table(rows=3, cols=5)
    headers = ["Product", "Opening", "Purchase", "Sale", "Closing"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    table.rows[1].cells[0].text = "LIV 52"
    table.rows[1].cells[1].text = "10"
    table.rows[1].cells[2].text = "0"
    table.rows[1].cells[3].text = "3"
    table.rows[1].cells[4].text = "7"
    table.rows[2].cells[0].text = "ABANA"
    table.rows[2].cells[1].text = "5"
    table.rows[2].cells[2].text = "2"
    table.rows[2].cells[3].text = "1"
    table.rows[2].cells[4].text = "6"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_image_only_docx_bytes() -> bytes:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    doc = Document()
    doc.add_paragraph("Scanned stock statement")
    img_buf = io.BytesIO()
    Image.new("RGB", (120, 80), color=(200, 200, 200)).save(img_buf, format="PNG")
    img_buf.seek(0)
    # Save temp png for add_picture
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp.write(img_buf.getvalue())
        tmp_path = tmp.name
    try:
        doc.add_picture(tmp_path, width=Inches(2))
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class WordLinuxFlagOffTests(unittest.TestCase):
    def setUp(self):
        ensure_gemini_offline_guard()
        self._env = {
            k: os.environ.get(k)
            for k in ("STOCK_WORD_LINUX", "STOCK_NATIVE_RESOLVER", "STOCK_VISION_TABLE")
        }
        os.environ.pop("STOCK_WORD_LINUX", None)
        os.environ["STOCK_NATIVE_RESOLVER"] = "off"
        os.environ.pop("STOCK_VISION_TABLE", None)

    def tearDown(self):
        for key, value in self._env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def test_flag_off(self):
        self.assertFalse(_stock_word_linux_enabled())

    def test_docx_byte_identical_path(self):
        data = _make_table_docx_bytes()
        a = _parse_word(data, "t.docx", ".docx")
        b = _parse_word(data, "t.docx", ".docx")
        self.assertEqual(
            [i.get("product_name") for i in a.get("line_items") or []],
            [i.get("product_name") for i in b.get("line_items") or []],
        )


class WordLinuxSofficeTests(unittest.TestCase):
    def setUp(self):
        ensure_gemini_offline_guard()
        self._env = os.environ.get("STOCK_WORD_LINUX")
        os.environ["STOCK_WORD_LINUX"] = "true"

    def tearDown(self):
        if self._env is None:
            os.environ.pop("STOCK_WORD_LINUX", None)
        else:
            os.environ["STOCK_WORD_LINUX"] = self._env

    def test_doc_converted_via_soffice(self):
        if not _soffice_path():
            self.skipTest("soffice/libreoffice not installed")
        docx_bytes = _make_table_docx_bytes()
        with tempfile.TemporaryDirectory() as td:
            docx_path = os.path.join(td, "sample.docx")
            doc_path = os.path.join(td, "sample.doc")
            with open(docx_path, "wb") as fh:
                fh.write(docx_bytes)
            # Convert docx → doc then parse as .doc
            subprocess.run(
                [
                    _soffice_path(),
                    "--headless",
                    "--norestore",
                    "--convert-to",
                    "doc",
                    "--outdir",
                    td,
                    docx_path,
                ],
                check=True,
                timeout=60,
                capture_output=True,
            )
            if not os.path.isfile(doc_path):
                self.skipTest("soffice did not produce .doc")
            with open(doc_path, "rb") as fh:
                doc_bytes = fh.read()
        result = _parse_word(doc_bytes, "sample.doc", ".doc")
        method = str(
            ((result.get("totals") or {}).get("extra") or {}).get("extraction_method")
            or ""
        )
        self.assertNotEqual(method, "legacy_doc_text")
        self.assertNotEqual(method, "legacy_doc_text_unverified")
        self.assertTrue(
            result.get("line_items") or method.startswith("doc_converted"),
            msg=f"method={method} items={len(result.get('line_items') or [])}",
        )

    def test_soffice_missing_legacy_unverified(self):
        fake_doc = b"\xd0\xcf\x11\xe0" + b"ALPHA MEDICAL STORES Product Opening Purchase Sale Closing LIV52 10 0 3 7" * 3
        with patch(
            "services.sales_statement_extractor._convert_doc_to_docx_bytes",
            return_value=None,
        ), patch("shutil.which", return_value=None):
            result = _parse_word(fake_doc, "x.doc", ".doc")
        method = ((result.get("totals") or {}).get("extra") or {}).get(
            "extraction_method"
        )
        self.assertEqual(method, "legacy_doc_text_unverified")
        quality = assess_extraction_quality(result, {"ext": ".doc"})
        self.assertLessEqual(int(quality.get("score") or 0), 50)
        self.assertTrue(quality.get("should_fallback"))

    def test_soffice_timeout_cleans_tmpdir(self):
        removed = []

        def _fake_rmtree(path, ignore_errors=False):
            removed.append(path)

        with patch("shutil.which", return_value="/usr/bin/soffice"), patch(
            "subprocess.run",
            side_effect=subprocess.TimeoutExpired(cmd="soffice", timeout=60),
        ), patch("shutil.rmtree", side_effect=_fake_rmtree), patch(
            "tempfile.mkdtemp", return_value="/tmp/stock_word_lo_testclean"
        ), patch("os.makedirs"), patch("builtins.open", MagicMock()), patch(
            "os.path.isfile", return_value=False
        ):
            # Also create the path so write doesn't fail oddly — open is mocked.
            out = _convert_doc_to_docx_bytes_linux(b"fake")
        self.assertIsNone(out)
        self.assertTrue(removed)


class WordImageOnlyVisionTests(unittest.TestCase):
    def setUp(self):
        ensure_gemini_offline_guard()
        reset_gemini_test_call_count()
        self._env = {
            k: os.environ.get(k)
            for k in ("STOCK_WORD_LINUX", "STOCK_VISION_TABLE", "STOCK_VISION_TABLE_TYPES")
        }
        os.environ["STOCK_WORD_LINUX"] = "true"

    def tearDown(self):
        for key, value in self._env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def test_image_only_calls_vision_when_on(self):
        os.environ["STOCK_VISION_TABLE"] = "true"
        os.environ["STOCK_VISION_TABLE_TYPES"] = "image"
        data = _make_image_only_docx_bytes()
        self.assertTrue(_docx_extract_media_images(data))
        called = []

        def _fake_vision(image_bytes, input_type, meta):
            called.append((input_type, meta.get("ext")))
            out = empty_result("img.png", "png")
            out["line_items"] = [
                {
                    "product_name": "LIV 52",
                    "opening_qty": 1.0,
                    "receipts_qty": 0.0,
                    "sales_qty": 0.0,
                    "closing_qty": 1.0,
                    "extra": {},
                }
            ]
            return {"status": "ok", "result": out, "gemini_calls": 1}

        result = empty_result("x.docx", "docx")
        with patch(
            "services.stock_vision_table.run_vision_table_path",
            side_effect=_fake_vision,
        ), patch(
            "services.stock_vision_table.vision_table_active_for",
            return_value=True,
        ):
            out = _docx_try_vision_images(result, data, "x.docx")
        self.assertEqual(len(called), 1)
        self.assertEqual(called[0][0], "image")
        self.assertEqual(
            ((out.get("totals") or {}).get("extra") or {}).get("extraction_method"),
            "docx_image_vision_table",
        )

    def test_image_only_skipped_when_vision_off(self):
        os.environ["STOCK_VISION_TABLE"] = "false"
        data = _make_image_only_docx_bytes()
        result = empty_result("x.docx", "docx")
        with patch(
            "services.stock_vision_table.vision_table_active_for",
            return_value=False,
        ), patch(
            "services.stock_vision_table.run_vision_table_path"
        ) as vision:
            out = _docx_try_vision_images(result, data, "x.docx")
        vision.assert_not_called()
        self.assertFalse(out.get("line_items"))


if __name__ == "__main__":
    unittest.main()
